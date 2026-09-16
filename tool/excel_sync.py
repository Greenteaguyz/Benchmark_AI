"""
Automated Real-Time Excel Synchronizer for Official Benchmark Runs.
Compliant with: KiTH AI Research Internship Project Plan (Page 3 & Page 6).

Maintains evaluation/master_scoring.xlsx:
- Populates newly generated official model responses into Column G (Complete Response)
  with verified answer on top and reasoning trace below, strictly wrapped.
- Populates real-time telemetry (Date/Time, Durations, Tokens, Peak VRAM GB, Status, Evidence).
- Strictly preserves human evaluator rubric scores (Columns H-K) and dynamic Excel formulas.
- Thread/Process-safe via cross-platform file locking (msvcrt / fcntl).
- Resilient to Microsoft Excel file locks: automatically queues records into
  evaluation/pending_sync.json if Excel has the file open, and flushes them
  as soon as the file is accessible.
"""

import os
import re
import io
import time
import json
import contextlib
from typing import Dict, Any, Optional, Tuple, List

try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, Border, Side
except ImportError:
    openpyxl = None

try:
    import msvcrt
except ImportError:
    msvcrt = None

try:
    import fcntl
except ImportError:
    fcntl = None


CURRENT_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_ROOT = os.path.dirname(CURRENT_DIR) if os.path.basename(CURRENT_DIR) in ("tool", "test_runs") else CURRENT_DIR
DEFAULT_EXCEL_PATH = os.path.join(PROJECT_ROOT, "evaluation", "master_scoring.xlsx")
RESPONSES_DIR = os.path.join(PROJECT_ROOT, "data", "responses")
PENDING_SYNC_PATH = os.path.join(PROJECT_ROOT, "evaluation", "pending_sync.json")

MODEL_ALIAS_MAP = {
    "phi4-mini-reasoning": "PHI",
    "phi4-mini-reasoning:latest": "PHI",
    "phi4": "PHI",
    "phi": "PHI",
    "deepseek-r1:7b": "DSR1",
    "deepseek-r1": "DSR1",
    "deepseek": "DSR1",
    "dsr1": "DSR1",
    "qwen3:8b": "QWEN",
    "qwen3": "QWEN",
    "qwen": "QWEN",
}


@contextlib.contextmanager
def file_lock(lock_path: str, timeout_sec: float = 10.0):
    """Guarantees process-safe concurrent access across scripts."""
    os.makedirs(os.path.dirname(lock_path), exist_ok=True)
    start_time = time.time()
    f = open(lock_path, "a+", encoding="utf-8")
    locked = False
    try:
        while True:
            try:
                if msvcrt:
                    f.seek(0)
                    if f.tell() == 0 and os.path.getsize(lock_path) == 0:
                        f.write(" ")
                        f.flush()
                    f.seek(0)
                    msvcrt.locking(f.fileno(), msvcrt.LK_NBLCK, 1)
                    locked = True
                    break
                elif fcntl:
                    fcntl.flock(f, fcntl.LOCK_EX | fcntl.LOCK_NB)
                    locked = True
                    break
                else:
                    locked = True
                    break
            except (OSError, IOError):
                if time.time() - start_time >= timeout_sec:
                    raise TimeoutError(f"Could not acquire lock on {lock_path} after {timeout_sec} seconds.")
                time.sleep(0.1)
        yield
    finally:
        if locked:
            try:
                if msvcrt:
                    f.seek(0)
                    msvcrt.locking(f.fileno(), msvcrt.LK_UNLCK, 1)
                elif fcntl:
                    fcntl.flock(f, fcntl.LOCK_UN)
            except (OSError, IOError):
                pass
        f.close()


def is_file_writable(filepath: str) -> bool:
    """Checks if the file is currently unlocked and writable by another process (e.g. MS Excel)."""
    if not os.path.exists(filepath):
        return True
    try:
        with open(filepath, "r+b"):
            return True
    except (IOError, PermissionError):
        return False


def parse_reasoning_and_answer(raw_text: str) -> Tuple[str, str]:
    """Separates <think> reasoning process from verified answer."""
    if not raw_text:
        return "", ""
    think_match = re.search(r"<think>(.*?)</think>", raw_text, flags=re.DOTALL)
    if think_match:
        thought = think_match.group(1).strip()
        answer = re.sub(r"<think>.*?</think>", "", raw_text, flags=re.DOTALL).strip()
        return thought, answer
    return "", raw_text.strip()


def format_complete_response(raw_text: str) -> str:
    """Formats model output with clean final answer on top and labeled reasoning below."""
    if not raw_text:
        return ""
    thought, answer = parse_reasoning_and_answer(raw_text)
    if thought and answer:
        return f"[FINAL ANSWER]\n{answer}\n\n[THINKING PROCESS (<think>)]\n{thought}"
    elif answer:
        return answer
    return raw_text.strip()


def normalize_qid(qid_raw: str) -> str:
    """Normalizes question ID into standard Q01-Q15 format."""
    s = str(qid_raw).strip().upper()
    digits = "".join(filter(str.isdigit, s))
    if digits:
        return f"Q{int(digits):02d}"
    return s


def normalize_model_alias(model_name: str, given_alias: Optional[str] = None) -> str:
    """Resolves standard model alias: PHI, DSR1, or QWEN."""
    if given_alias and given_alias.strip().upper() in ("PHI", "DSR1", "QWEN"):
        return given_alias.strip().upper()
    cleaned = str(model_name).strip().lower()
    return MODEL_ALIAS_MAP.get(cleaned, cleaned.upper()[:4])


def load_pending_queue(pending_path: str = PENDING_SYNC_PATH) -> Dict[str, Any]:
    """Loads any queued updates that couldn't be written because Excel had the file open."""
    if os.path.exists(pending_path):
        try:
            with open(pending_path, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            return {}
    return {}


def save_pending_queue(queue: Dict[str, Any], pending_path: str = PENDING_SYNC_PATH) -> None:
    """Persists queued updates."""
    os.makedirs(os.path.dirname(pending_path), exist_ok=True)
    if not queue:
        if os.path.exists(pending_path):
            try:
                os.remove(pending_path)
            except OSError:
                pass
        return
    with open(pending_path, "w", encoding="utf-8") as f:
        json.dump(queue, f, indent=2)


def get_pending_sync_count() -> int:
    """Returns how many records are currently queued awaiting Excel closure."""
    q = load_pending_queue()
    return len(q)


def apply_single_record_to_sheet(ws, target_resp_id: str, qid: str, alias: str, model_name: str, record: Dict[str, Any], response_text: str) -> bool:
    """Updates a single row in ws (Scoring Register). Returns True if row found and updated."""
    # Locate target row (Rows 2 to ws.max_row)
    target_row_idx = None
    for r in range(2, ws.max_row + 1):
        cell_val = str(ws.cell(row=r, column=1).value or "").strip().upper()
        if cell_val == target_resp_id:
            target_row_idx = r
            break

    # Fallback match on QID (Col 2) and Model (Col 3)
    if not target_row_idx:
        for r in range(2, ws.max_row + 1):
            col_b = str(ws.cell(row=r, column=2).value or "").strip().upper()
            col_c = str(ws.cell(row=r, column=3).value or "").strip().lower()
            if col_b == qid and (alias.lower() in col_c or model_name.lower() in col_c):
                target_row_idx = r
                break

    if not target_row_idx:
        return False

    # Common styling definitions
    cell_font = Font(name="Segoe UI", size=10)
    thin_border = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'),
        bottom=Side(style='thin', color='E5E7EB')
    )
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Parse telemetry fields safely
    ts = record.get("timestamp") or time.strftime("%Y-%m-%d %H:%M:%S")

    try:
        total_dur = round(float(record.get("total_duration_s")), 3)
    except (ValueError, TypeError):
        total_dur = None

    try:
        gen_dur = round(float(record.get("generation_duration_s")), 3)
    except (ValueError, TypeError):
        gen_dur = None

    try:
        out_toks = int(record.get("output_tokens"))
    except (ValueError, TypeError):
        out_toks = None

    # VRAM: convert MB to GB if > 100
    raw_vram = record.get("peak_vram_gb") or record.get("peak_vram_mb")
    peak_vram = None
    if raw_vram is not None and str(raw_vram).strip() != "":
        try:
            v_num = float(raw_vram)
            peak_vram = round(v_num / 1024, 2) if v_num > 100 else round(v_num, 2)
        except (ValueError, TypeError):
            peak_vram = None

    raw_status = str(record.get("completion_status", "Completed")).strip()
    comp_status = "Completed"
    status_note = ""
    if raw_status.lower() == "length":
        comp_status = "Completed"
        status_note = "Reached max tokens limit (num_predict=1024)"
    elif raw_status in ("Completed", "Timeout", "OOM", "Error"):
        comp_status = raw_status

    evidence_path = f"data/responses/{target_resp_id}.json"
    formatted_response = format_complete_response(response_text or "")

    # Update Column 7 (G): Complete Response
    cell_g = ws.cell(row=target_row_idx, column=7)
    cell_g.value = formatted_response
    cell_g.font = cell_font
    cell_g.alignment = align_left
    cell_g.border = thin_border

    # Columns 8, 9, 10, 11 (H, I, J, K): DO NOT OVERWRITE HUMAN EVALUATION SCORES!
    for col_idx in (8, 9, 10, 11):
        sc_cell = ws.cell(row=target_row_idx, column=col_idx)
        sc_cell.font = cell_font
        sc_cell.alignment = align_center
        sc_cell.border = thin_border

    # Column 12 (L): Formula =IF(COUNT(H:K)>0, SUM(H:K), "")
    cell_l = ws.cell(row=target_row_idx, column=12)
    total_formula = f'=IF(COUNT(H{target_row_idx}:K{target_row_idx})>0, SUM(H{target_row_idx}:K{target_row_idx}), "")'
    if not cell_l.value or str(cell_l.value).strip() == "":
        cell_l.value = total_formula
    cell_l.font = cell_font
    cell_l.alignment = align_center
    cell_l.border = thin_border

    # Column 13 (M): Date & Start Time
    cell_m = ws.cell(row=target_row_idx, column=13)
    cell_m.value = ts
    cell_m.font = cell_font
    cell_m.alignment = align_center
    cell_m.border = thin_border

    # Column 14 (N): Total Duration (s)
    cell_n = ws.cell(row=target_row_idx, column=14)
    cell_n.value = total_dur
    cell_n.font = cell_font
    cell_n.alignment = align_center
    cell_n.border = thin_border

    # Column 15 (O): Generation Duration (s)
    cell_o = ws.cell(row=target_row_idx, column=15)
    cell_o.value = gen_dur
    cell_o.font = cell_font
    cell_o.alignment = align_center
    cell_o.border = thin_border

    # Column 16 (P): Output Tokens
    cell_p = ws.cell(row=target_row_idx, column=16)
    cell_p.value = out_toks
    cell_p.font = cell_font
    cell_p.alignment = align_center
    cell_p.border = thin_border

    # Column 17 (Q): Tokens/sec Formula
    cell_q = ws.cell(row=target_row_idx, column=17)
    tps_formula = (
        f'=IF(AND(ISNUMBER(O{target_row_idx}), O{target_row_idx}>0, ISNUMBER(P{target_row_idx})), '
        f'ROUND(P{target_row_idx}/O{target_row_idx}, 2), '
        f'IF(AND(ISNUMBER(N{target_row_idx}), N{target_row_idx}>0, ISNUMBER(P{target_row_idx})), '
        f'ROUND(P{target_row_idx}/N{target_row_idx}, 2), ""))'
    )
    if not cell_q.value or str(cell_q.value).strip() == "":
        cell_q.value = tps_formula
    cell_q.font = cell_font
    cell_q.alignment = align_center
    cell_q.border = thin_border

    # Column 18 (R): Peak VRAM (GB)
    cell_r = ws.cell(row=target_row_idx, column=18)
    cell_r.value = peak_vram
    cell_r.font = cell_font
    cell_r.alignment = align_center
    cell_r.border = thin_border

    # Column 19 (S): Completion Status
    cell_s = ws.cell(row=target_row_idx, column=19)
    cell_s.value = comp_status
    cell_s.font = cell_font
    cell_s.alignment = align_center
    cell_s.border = thin_border

    # Column 20 (T): Verification Status
    cell_t = ws.cell(row=target_row_idx, column=20)
    if not cell_t.value:
        cell_t.value = "Pending"
    cell_t.font = cell_font
    cell_t.alignment = align_center
    cell_t.border = thin_border

    # Column 21 (U): Evidence Reference
    cell_u = ws.cell(row=target_row_idx, column=21)
    cell_u.value = evidence_path
    cell_u.font = cell_font
    cell_u.alignment = align_left
    cell_u.border = thin_border

    # Column 22 (V): Evaluator Notes
    cell_v = ws.cell(row=target_row_idx, column=22)
    if not cell_v.value and status_note:
        cell_v.value = status_note
    cell_v.font = cell_font
    cell_v.alignment = align_left
    cell_v.border = thin_border

    return True


def _safe_print(text: str) -> None:
    """Prints text safely without throwing UnicodeEncodeError on Windows cp1252 console."""
    try:
        print(text)
    except UnicodeEncodeError:
        safe_text = text.encode("ascii", errors="backslashreplace").decode("ascii")
        print(safe_text)


def flush_pending_syncs(workbook_path: Optional[str] = None) -> Tuple[bool, str]:
    """Flushes queued updates to master_scoring.xlsx if the file is unlocked and writable."""
    target_path = os.path.abspath(workbook_path or DEFAULT_EXCEL_PATH)
    queue = load_pending_queue()
    if not queue:
        return True, "No pending records in queue."

    if not is_file_writable(target_path):
        return False, f"Cannot flush: {os.path.basename(target_path)} is still locked by Microsoft Excel. Please close it first."

    lock_file = f"{target_path}.lock"
    with file_lock(lock_file):
        try:
            wb = openpyxl.load_workbook(target_path)
        except Exception as e:
            return False, f"Failed to open workbook: {e}"

        if "Scoring Register" not in wb.sheetnames:
            wb.close()
            return False, "Scoring Register sheet not found."

        ws = wb["Scoring Register"]
        flushed_keys = []
        for resp_id, item in list(queue.items()):
            rec = item["record"]
            resp_txt = item.get("response_text", "")
            qid = normalize_qid(rec.get("question_id", ""))
            alias = normalize_model_alias(rec.get("model", ""), rec.get("model_alias"))
            ok = apply_single_record_to_sheet(ws, resp_id, qid, alias, rec.get("model", ""), rec, resp_txt)
            if ok:
                flushed_keys.append(resp_id)

        try:
            wb.save(target_path)
            wb.close()
            for k in flushed_keys:
                queue.pop(k, None)
            save_pending_queue(queue)
            msg = f"Successfully flushed {len(flushed_keys)} queued record(s) into {os.path.basename(target_path)}."
            _safe_print(f"[Excel Sync] [OK] {msg}")
            return True, msg
        except PermissionError:
            wb.close()
            return False, f"{os.path.basename(target_path)} is locked by Microsoft Excel. Queued records retained."
        except Exception as e:
            wb.close()
            return False, f"Failed to save workbook: {e}"


def sync_official_benchmark_to_excel(
    record: Dict[str, Any],
    response_text: Optional[str] = None,
    workbook_path: Optional[str] = None,
) -> Tuple[bool, str]:
    """
    Synchronizes an official benchmark result directly into evaluation/master_scoring.xlsx.
    
    If the Excel file is open in Microsoft Excel, queues the record in pending_sync.json
    so no data is lost, and attempts to flush all pending records on the next available write.
    
    Returns:
        Tuple[bool, str]: (Success, User-friendly status message)
    """
    if openpyxl is None:
        msg = "Warning: openpyxl is not installed. Excel sync skipped."
        _safe_print(f"[Excel Sync] {msg}")
        return False, msg

    target_path = os.path.abspath(workbook_path or DEFAULT_EXCEL_PATH)
    lock_file = f"{target_path}.lock"

    # Identify QID and Model Alias
    qid = normalize_qid(record.get("question_id", ""))
    model_name = record.get("model", "")
    alias = normalize_model_alias(model_name, record.get("model_alias"))
    target_resp_id = f"{qid}_{alias}"

    # Only official Q01 to Q15 runs map to the master scoring sheet
    try:
        q_num = int(qid.replace("Q", ""))
        if not (1 <= q_num <= 15):
            msg = f"Notice: {qid} is outside official Q01-Q15 range. Skipped master sync."
            _safe_print(f"[Excel Sync] {msg}")
            return False, msg
    except ValueError:
        msg = f"Notice: Non-standard Question ID '{qid}'. Skipped master sync."
        _safe_print(f"[Excel Sync] {msg}")
        return False, msg

    if alias not in ("PHI", "DSR1", "QWEN"):
        msg = f"Notice: Model alias '{alias}' is not one of PHI, DSR1, QWEN. Skipped master sync."
        _safe_print(f"[Excel Sync] {msg}")
        return False, msg

    # Retrieve response text if not provided
    if not response_text:
        evidence_file = record.get("evidence_file")
        if evidence_file and os.path.exists(evidence_file):
            try:
                with open(evidence_file, "r", encoding="utf-8") as f:
                    data = json.load(f)
                response_text = data.get("response", "")
            except Exception:
                response_text = ""
        if not response_text:
            resp_json_path = os.path.join(RESPONSES_DIR, f"{target_resp_id}.json")
            if os.path.exists(resp_json_path):
                try:
                    with open(resp_json_path, "r", encoding="utf-8") as f:
                        data = json.load(f)
                    response_text = data.get("response", "")
                except Exception:
                    response_text = ""

    # Check if target workbook exists. If not, generate via build_workbook()
    if not os.path.exists(target_path):
        _safe_print(f"[Excel Sync] {target_path} does not exist. Initializing master workbook...")
        try:
            from tool.generate_scoring_sheet import build_workbook
            build_workbook()
        except Exception as e:
            msg = f"Failed to auto-generate workbook: {e}"
            _safe_print(f"[Excel Sync] {msg}")
            return False, msg

    # Check if file is currently open and locked by Microsoft Excel
    if not is_file_writable(target_path):
        # File is locked in Excel: queue it into pending_sync.json
        queue = load_pending_queue()
        queue[target_resp_id] = {
            "record": record,
            "response_text": response_text or "",
            "queued_at": time.strftime("%Y-%m-%d %H:%M:%S")
        }
        save_pending_queue(queue)
        msg = f"[!] master_scoring.xlsx is open in Microsoft Excel. Telemetry for {target_resp_id} queued safely and will sync once Excel is closed."
        _safe_print(f"[Excel Sync] {msg}")
        return False, msg

    # File is writable: acquire cross-process lock and apply
    with file_lock(lock_file):
        # Also flush any previously queued updates
        queue = load_pending_queue()
        queue[target_resp_id] = {
            "record": record,
            "response_text": response_text or "",
            "queued_at": time.strftime("%Y-%m-%d %H:%M:%S")
        }

        try:
            wb = openpyxl.load_workbook(target_path)
        except Exception as e:
            msg = f"Error loading workbook '{target_path}': {e}"
            _safe_print(f"[Excel Sync] {msg}")
            return False, msg

        if "Scoring Register" not in wb.sheetnames:
            msg = f"'Scoring Register' sheet not found in '{target_path}'."
            _safe_print(f"[Excel Sync] {msg}")
            wb.close()
            return False, msg

        ws = wb["Scoring Register"]
        flushed_keys = []
        for resp_id_item, item in list(queue.items()):
            rec = item["record"]
            resp_txt = item.get("response_text", "")
            q_id = normalize_qid(rec.get("question_id", ""))
            m_alias = normalize_model_alias(rec.get("model", ""), rec.get("model_alias"))
            ok = apply_single_record_to_sheet(ws, resp_id_item, q_id, m_alias, rec.get("model", ""), rec, resp_txt)
            if ok:
                flushed_keys.append(resp_id_item)

        try:
            wb.save(target_path)
            wb.close()
            for k in flushed_keys:
                queue.pop(k, None)
            save_pending_queue(queue)
            msg = f"[OK] Successfully synced {target_resp_id} to master_scoring.xlsx."
            _safe_print(f"[Excel Sync] {msg}")
            return True, msg
        except PermissionError:
            wb.close()
            # Still locked unexpectedly, save queue
            save_pending_queue(queue)
            msg = f"[!] master_scoring.xlsx is locked by Microsoft Excel. Queued {target_resp_id} in pending_sync.json."
            _safe_print(f"[Excel Sync] {msg}")
            return False, msg
        except Exception as e:
            wb.close()
            msg = f"Failed to save workbook: {e}"
            _safe_print(f"[Excel Sync] {msg}")
            return False, msg
